"""The manuscript builder renders headings, lists, tables and inline markup, and counts placeholders."""

import json

import docx
import pytest

from scripts.build_docx import convert, main

SAMPLE = """# Title

> **Status: DRAFT.** Editorial banner, never submitted.

## Methods

The score is *S* = *w*<sub>r</sub>*R*, with **bold** text and `code`.

- first bullet
- second bullet
  - nested bullet

| Policy | Value |
|---|---|
| DARS | 0.5 |
| LRU | 0.4 |

## Results

Recall was ⟦value⟧ on the test split.

⟦author note (dev) — delete me before sending.⟧

## Discussion

Two words.
"""


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def test_headings_and_inline_markup():
    doc, stats = convert(SAMPLE)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Title", "Methods", "Results", "Discussion"]
    body = next(p for p in doc.paragraphs if p.text.startswith("The score is"))
    assert any(r.italic and r.text == "S" for r in body.runs)
    assert any(r.font.subscript and r.text == "r" for r in body.runs)
    assert any(r.bold and r.text == "bold" for r in body.runs)
    assert any(r.font.name == "Consolas" and r.text == "code" for r in body.runs)


def test_lists_and_table():
    doc, _ = convert(SAMPLE)
    styles = [p.style.name for p in doc.paragraphs if p.text.endswith("bullet")]
    assert styles[:2] == ["List Bullet", "List Bullet"]
    assert styles[2].startswith("List Bullet")            # nested level
    table = doc.tables[0]
    assert len(table.rows) == 3 and len(table.columns) == 2
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert [c.text for c in table.rows[1].cells] == ["DARS", "0.5"]


def test_placeholders_are_counted_and_highlighted():
    doc, stats = convert(SAMPLE)
    assert stats["placeholders"] == ["⟦value⟧"]           # the author note was dropped
    para = next(p for p in doc.paragraphs if "⟦value⟧" in p.text)
    run = next(r for r in para.runs if r.text == "⟦value⟧")
    assert run.font.highlight_color == docx.enum.text.WD_COLOR_INDEX.YELLOW


def test_editorial_notes_are_dropped_by_default_and_kept_on_request():
    assert not any("Editorial banner" in t for t in _texts(convert(SAMPLE)[0]))
    assert not any("author note" in t for t in _texts(convert(SAMPLE)[0]))
    kept = _texts(convert(SAMPLE, drop_notes=False)[0])
    assert any("Editorial banner" in t for t in kept)


def test_word_count_covers_only_the_narrative_sections():
    _, stats = convert(SAMPLE)
    assert stats["main_text_words"] == len("Recall was ⟦value⟧ on the test split.".split()) + 2
    assert "methods" in stats["words_by_section"]         # counted, but not in the budget


REFS = {"_comment": "ignored", "RAG": "Lewis, P. et al. Retrieval-augmented generation. NeurIPS 33 (2020).",
        "BM25": "Robertson, S. & Zaragoza, H. The probabilistic relevance framework. FnTIR 3, 333-389 (2009)."}

CITED = """## Introduction

Retrieval augmentation⟦ref:RAG⟧ and lexical search⟦ref:BM25⟧ matter; RAG again⟦ref:RAG⟧ and both⟦ref:BM25,RAG⟧.

## References

1. A stale hand-numbered list that must never be copied into the submission.
"""


def test_citations_are_numbered_in_order_of_first_use():
    doc, stats = convert(CITED, refs=REFS)
    para = next(p for p in doc.paragraphs if p.text.startswith("Retrieval augmentation"))
    assert [r.text for r in para.runs if r.font.superscript] == ["1", "2", "1", "2,1"]
    assert stats["citations"] == 2 and stats["unknown_refs"] == []


def test_reference_list_comes_from_the_verified_bibliography():
    doc, _ = convert(CITED, refs=REFS)
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("1. Lewis, P. et al.") for t in texts)
    assert any(t.startswith("2. Robertson, S. & Zaragoza, H.") for t in texts)
    assert not any("stale hand-numbered" in t for t in texts)
    assert sum(t == "References" for t in texts) == 1


def test_unknown_key_is_reported_and_blocks_strict(tmp_path):
    _, stats = convert("Text⟦ref:NotInBibliography⟧.\n", refs=REFS)
    assert stats["unknown_refs"] == ["NotInBibliography"]
    src, refs, out = tmp_path / "m.md", tmp_path / "r.json", tmp_path / "m.docx"
    src.write_text("Text⟦ref:NotInBibliography⟧.\n", encoding="utf-8")
    refs.write_text(json.dumps(REFS), encoding="utf-8")
    assert main(["--in", str(src), "--out", str(out), "--refs", str(refs), "--strict"]) == 1
    assert not out.exists()


def test_strict_mode_refuses_to_write_a_draft(tmp_path):
    src = tmp_path / "m.md"
    src.write_text(SAMPLE, encoding="utf-8")
    out = tmp_path / "m.docx"
    assert main(["--in", str(src), "--out", str(out), "--strict"]) == 1
    assert not out.exists()
    assert main(["--in", str(src), "--out", str(out)]) == 0
    assert out.stat().st_size > 0
    docx.Document(str(out))                               # the file opens


WRAPPED = """## Results

### A subsection

This sentence is wrapped
across three lines
of Markdown.
- a bullet whose text
  continues here
  - a nested bullet
    that also wraps
- a second bullet

  | Col | Val |
  |---|---|
  | a | 1 |

```
command --flag one
command --flag two
```

## Methods

Not counted.
"""


def test_hard_wrapped_lines_join_into_one_paragraph_or_list_item():
    doc, stats = convert(WRAPPED)
    texts = [p.text for p in doc.paragraphs if p.text]
    assert "This sentence is wrapped across three lines of Markdown." in texts
    assert "a bullet whose text continues here" in texts
    assert "a nested bullet that also wraps" in texts
    assert "a second bullet" in texts
    assert not any(t.startswith(("across", "continues", "that also")) for t in texts)
    assert len(doc.tables) == 1 and [c.text for c in doc.tables[0].rows[1].cells] == ["a", "1"]
    code = [p for p in doc.paragraphs if p.text.startswith("command")]
    assert len(code) == 2 and code[0].runs[0].font.name == "Consolas"
    assert not any("```" in t for t in texts)


def test_word_count_includes_subsections_of_narrative_sections():
    _, stats = convert(WRAPPED)
    narrative = "This sentence is wrapped across three lines of Markdown."
    bullets = "a bullet whose text continues here a nested bullet that also wraps a second bullet"
    code = "command --flag one command --flag two"
    heading_free = len(narrative.split()) + len(bullets.split()) + len(code.split())
    assert stats["main_text_words"] == heading_free
    assert stats["words_by_section"]["methods"] == 2
