"""Build the submission .docx from the Markdown manuscript (pandoc is not available here).

Produces a *Scientific Reports* Article layout: Times New Roman 12 pt, double spacing,
continuous line numbers and page numbers, headings, lists, tables and inline markup
(``**bold**``, ``*italic*``, ``` `code` ```, ``<sub>``/``<sup>``).

Unfilled ⟦…⟧ placeholders are highlighted yellow and counted, so a draft can never be
mistaken for a finished file.  ``--strict`` refuses to write while any remain.

    python scripts/build_docx.py --in manuscript.md --out manuscript.docx [--strict]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
SIZE = Pt(12)
PLACEHOLDER = re.compile(r"⟦(?!ref:)[^⟧]*⟧")
CITATION = re.compile(r"⟦ref:([A-Za-z0-9_,\-]+)⟧")
# bold / italic / code / sub / sup / placeholder, in one pass
INLINE = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`|<sub>.*?</sub>|<sup>.*?</sup>|⟦[^⟧]*⟧)",
    re.DOTALL,
)
NO_REFERENCE_SECTIONS = ("references",)      # written from the bibliography, not the Markdown
# the word budget covers the narrative sections only
MAIN_TEXT_SECTIONS = ("introduction", "results", "discussion")


def _set_line_numbers(section) -> None:
    """Continuous line numbers, which reviewers expect on a submitted manuscript."""
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    section._sectPr.append(ln)


def _add_page_number_footer(section) -> None:
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    for element, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                                 ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
                                 ("w:fldChar", {"w:fldCharType": "end"}, None)):
        node = OxmlElement(element)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text is not None:
            node.text = text
        run._r.append(node)


def new_document() -> "docx.document.Document":
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    fmt = style.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_after = Pt(0)
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    for side in ("top", "bottom", "left", "right"):
        setattr(section, f"{side}_margin", Inches(1))
    _set_line_numbers(section)
    _add_page_number_footer(section)
    return doc


class Citations:
    """Assigns reference numbers in order of first citation, as Nature-style journals require.

    Numbering by hand is what produced the wrong references in the submitted version, so
    every ``⟦ref:Key⟧`` marker is resolved against the verified bibliography instead.
    """

    def __init__(self, refs: Optional[Dict[str, str]] = None):
        self.refs = {k: v for k, v in (refs or {}).items() if not k.startswith("_")}
        self.numbers: Dict[str, int] = {}
        self.unknown: List[str] = []

    def number(self, key: str) -> int:
        if key not in self.numbers:
            self.numbers[key] = len(self.numbers) + 1
            if key not in self.refs:
                self.unknown.append(key)
        return self.numbers[key]

    def marker(self, keys: str) -> str:
        return ",".join(str(self.number(k.strip())) for k in keys.split(",") if k.strip())

    def ordered(self) -> List[Tuple[int, str]]:
        return [(n, self.refs.get(k, f"MISSING REFERENCE: {k}"))
                for k, n in sorted(self.numbers.items(), key=lambda kv: kv[1])]


def add_runs(paragraph, text: str, citations: Optional[Citations] = None) -> None:
    """Write ``text`` into ``paragraph``, honouring the inline markup."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        cite = CITATION.fullmatch(piece)
        if cite is not None:
            run = paragraph.add_run(citations.marker(cite.group(1)) if citations else cite.group(1))
            run.font.superscript = True
        elif piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        elif piece.startswith("<sub>"):
            paragraph.add_run(piece[5:-6]).font.subscript = True
        elif piece.startswith("<sup>"):
            paragraph.add_run(piece[5:-6]).font.superscript = True
        elif PLACEHOLDER.fullmatch(piece):
            run = paragraph.add_run(piece)
            run.font.highlight_color = docx.enum.text.WD_COLOR_INDEX.YELLOW
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _split_row(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: List[str], citations: Optional[Citations] = None) -> None:
    header = _split_row(rows[0])
    body = [_split_row(r) for r in rows[2:]]          # rows[1] is the --- separator
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        add_runs(cell.paragraphs[0], text, citations)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for line in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, line[: len(header)]):
            add_runs(cell.paragraphs[0], text, citations)
    doc.add_paragraph()


def _starts_block(line: str) -> bool:
    """True if ``line`` opens a new Markdown block instead of continuing the previous paragraph."""
    s = line.strip()
    return (s.startswith(("#", ">", "```", "|")) or s == "---"
            or re.match(r"^\s*([-*]|\d+\.)\s+", line) is not None)


def convert(md: str, *, drop_notes: bool = True,
            refs: Optional[Dict[str, str]] = None) -> Tuple["docx.document.Document", Dict[str, Any]]:
    """Render Markdown into a document; also return counts the author needs."""
    doc = new_document()
    citations = Citations(refs)
    lines = md.splitlines()
    section = top_section = ""
    words: Dict[str, int] = {}
    placeholders: List[str] = []
    skip_section = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):                 # fenced block: one monospace paragraph per line, no markup
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                placeholders.extend(PLACEHOLDER.findall(lines[i]))
                if not skip_section:
                    run = doc.add_paragraph().add_run(lines[i].rstrip())
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    words[top_section] = words.get(top_section, 0) + len(lines[i].split())
                i += 1
            i += 1
            continue
        if drop_notes and "⟦author note" in stripped:
            i += 1
            continue
        placeholders.extend(PLACEHOLDER.findall(stripped))
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith(">"):                   # editorial banner, never submitted
            if not drop_notes:
                add_runs(doc.add_paragraph(), stripped.lstrip("> ").strip())
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            section = title.lower()
            if level <= 2:
                top_section = section
            skip_section = section.startswith(NO_REFERENCE_SECTIONS)
            if not skip_section:
                add_runs(doc.add_heading("", level=min(level, 4)), title, citations)
            i += 1
            continue
        if skip_section:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, citations)
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        number = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        match = bullet or number
        # Markdown hard-wraps: following lines continue this paragraph or list item until a blank line or a new block;
        # a list item's continuation lines are indented past its marker.
        text, indent_chars = (match.group(2), len(match.group(1))) if match else (stripped, -1)
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if (not nxt.strip() or _starts_block(nxt)
                    or (match and len(nxt) - len(nxt.lstrip()) <= indent_chars)):
                break
            if drop_notes and "⟦author note" in nxt:
                break
            placeholders.extend(PLACEHOLDER.findall(nxt))
            text = f"{text} {nxt.strip()}"
            i += 1
        if match:
            indent = len(match.group(1)) // 2
            style = "List Bullet" if bullet else "List Number"
            style = style if indent == 0 else f"{style} {min(indent + 1, 3)}"
            para = doc.add_paragraph(style=style)
        else:
            para = doc.add_paragraph()
        add_runs(para, text, citations)
        words[top_section] = words.get(top_section, 0) + len(text.split())

    if citations.numbers:
        add_runs(doc.add_heading("", level=2), "References")
        for number, entry in citations.ordered():
            para = doc.add_paragraph()
            para.add_run(f"{number}. ")
            add_runs(para, entry)

    main = sum(n for s, n in words.items() if any(s.startswith(k) for k in MAIN_TEXT_SECTIONS))
    return doc, {"placeholders": placeholders, "words_by_section": words, "main_text_words": main,
                 "citations": len(citations.numbers), "unknown_refs": citations.unknown}


def main(argv: Optional[List[str]] = None) -> int:
    p = argparse.ArgumentParser(description="Markdown manuscript to Scientific Reports .docx")
    p.add_argument("--in", dest="src", required=True)
    p.add_argument("--out", dest="dest", required=True)
    p.add_argument("--keep-notes", action="store_true", help="keep editorial banners and author notes")
    p.add_argument("--refs", help="JSON of citation key to verified reference string")
    p.add_argument("--strict", action="store_true",
                   help="fail if any ⟦…⟧ placeholder or unknown citation key remains")
    args = p.parse_args(argv)

    refs = json.loads(Path(args.refs).read_text(encoding="utf-8")) if args.refs else None
    doc, stats = convert(Path(args.src).read_text(encoding="utf-8"),
                         drop_notes=not args.keep_notes, refs=refs)
    n = len(stats["placeholders"])
    unknown = stats["unknown_refs"]
    if (n or unknown) and args.strict:
        print(f"{n} unfilled placeholder(s), {len(unknown)} unknown citation key(s); "
              f"refusing to write {args.dest}", file=sys.stderr)
        for item in stats["placeholders"][:20] + [f"ref:{k}" for k in unknown[:20]]:
            print(f"  {item}", file=sys.stderr)
        return 1
    doc.save(args.dest)
    print(f"wrote {args.dest}")
    print(f"  main text (Introduction+Results+Discussion): {stats['main_text_words']} words (limit 4,500)")
    print(f"  unfilled placeholders: {n}")
    print(f"  references cited: {stats['citations']}" + (f"  UNKNOWN KEYS: {unknown}" if unknown else ""))
    if stats["main_text_words"] > 4500:
        print("  WARNING: over the Scientific Reports word limit", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
